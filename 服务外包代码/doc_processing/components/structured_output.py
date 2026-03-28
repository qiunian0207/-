from __future__ import annotations

import json
import os
import re
import unicodedata
from datetime import datetime
from typing import Any

import pandas as pd

try:
    from docx import Document
    from docx.shared import Inches
except ImportError:
    Document = None
    Inches = None

try:
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    Font = None
    PatternFill = None
    get_column_letter = None

try:
    from config.config import OUTPUT_DIR
    from components.chart_renderer import ChartRenderer
except ImportError:
    from doc_processing.config.config import OUTPUT_DIR
    from doc_processing.components.chart_renderer import ChartRenderer


DOMAIN_LABELS = {
    "general": "通用",
    "government_statistics": "政府统计",
    "environment_monitoring": "环境监测",
    "healthcare": "医疗健康",
    "finance": "金融财务",
    "manufacturing": "制造业",
    "education": "教育科研",
    "agriculture": "农业",
    "transportation": "交通物流",
    "energy": "能源",
    "retail_consumption": "零售消费",
    "human_resources": "人力资源",
    "legal_compliance": "法务合规",
    "technology_operations": "技术运维",
}
DOCUMENT_TYPE_LABELS = {
    "tabular_dataset": "表格数据集",
    "report_document": "报告文档",
    "key_value_document": "字段文档",
    "plain_text_document": "纯文本",
}
HEADER_FILL = "D9EAF7"


class StructuredOutput:
    def __init__(self) -> None:
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        self.chart_renderer = ChartRenderer(output_dir=OUTPUT_DIR)

    def generate(self, data: Any, output_format: str) -> dict[str, Any]:
        try:
            payload = self._normalize_payload(data)
            if output_format == "json":
                return self._generate_json(payload)
            if output_format == "excel":
                return self._generate_excel(payload)
            if output_format == "word":
                return self._generate_word(payload)
            if output_format == "text":
                return self._generate_text(payload)
            if output_format == "pdf":
                return self._generate_pdf(payload)
            raise ValueError(f"Unsupported output format: {output_format}")
        except Exception as exc:
            return {"format": output_format, "path": "", "error": str(exc)}

    def _generate_json(self, payload: dict[str, Any]) -> dict[str, Any]:
        output_path = self._build_output_path(payload, "json")
        with open(output_path, "w", encoding="utf-8") as file:
            json.dump(payload, file, ensure_ascii=False, indent=2)
        return self._build_result("json", output_path)

    def _generate_text(self, payload: dict[str, Any]) -> dict[str, Any]:
        output_path = self._build_output_path(payload, "txt")
        content = str(payload.get("natural_language", {}).get("content") or "").strip()
        if not content:
            content = self._fallback_report_text(payload)
        with open(output_path, "w", encoding="utf-8") as file:
            file.write(content)
        return self._build_result("text", output_path)

    def _generate_pdf(self, payload: dict[str, Any]) -> dict[str, Any]:
        output_path = self._build_output_path(payload, "pdf")
        pdf_text = self._pdf_source_text(payload)
        pdf_lines = self._pdf_lines(pdf_text)
        pdf_bytes = self._build_pdf_document(pdf_lines)
        with open(output_path, "wb") as file:
            file.write(pdf_bytes)
        return self._build_result("pdf", output_path)

    def _generate_excel(self, payload: dict[str, Any]) -> dict[str, Any]:
        output_path = self._build_output_path(payload, "xlsx")
        structured_data = payload.get("structured_data", {})
        natural_language = payload.get("natural_language", {})
        canonical_data = payload.get("canonical_data", {})
        datasets = structured_data.get("datasets", []) if isinstance(structured_data.get("datasets"), list) else []
        records = structured_data.get("records", []) if isinstance(structured_data.get("records"), list) else []
        sections = structured_data.get("sections", []) if isinstance(structured_data.get("sections"), list) else []
        indicators = structured_data.get("indicators", []) if isinstance(structured_data.get("indicators"), list) else []

        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            self._write_sheet(writer, "overview", pd.DataFrame(self._overview_rows(payload)))
            self._write_sheet(writer, "sheet_index", pd.DataFrame(self._detail_index_rows(payload)))

            pipeline_rows = self._pipeline_plan_rows(payload)
            if pipeline_rows:
                self._write_sheet(writer, "pipeline_plan", pd.DataFrame(pipeline_rows))

            report_fact_rows = self._report_fact_rows(payload.get("report_facts", {}))
            if report_fact_rows:
                self._write_sheet(writer, "report_facts", pd.DataFrame(report_fact_rows))

            report_topic_rows = self._report_topic_rows(payload.get("report_facts", {}))
            if report_topic_rows:
                self._write_sheet(writer, "report_topics", pd.DataFrame(report_topic_rows))

            report_plan_rows = self._report_plan_rows(payload.get("report_plan", {}))
            if report_plan_rows:
                self._write_sheet(writer, "report_plan", pd.DataFrame(report_plan_rows))

            chart_spec_rows = self._chart_spec_rows(payload.get("chart_specs", []))
            if chart_spec_rows:
                self._write_sheet(writer, "chart_specs", pd.DataFrame(chart_spec_rows))

            chart_output_rows = self._chart_output_rows(payload.get("chart_outputs", []))
            if chart_output_rows:
                self._write_sheet(writer, "chart_outputs", pd.DataFrame(chart_output_rows))

            paragraphs = natural_language.get("paragraphs", []) if isinstance(natural_language.get("paragraphs"), list) else []
            if paragraphs:
                self._write_sheet(writer, "narrative", pd.DataFrame([{"段落": item} for item in paragraphs]))

            metadata_rows = self._metadata_rows(structured_data, canonical_data)
            if metadata_rows:
                self._write_sheet(writer, "metadata", pd.DataFrame(metadata_rows))

            if records:
                self._write_sheet(writer, "structured_data", pd.DataFrame(self._record_rows(records, structured_data, canonical_data)))
            if sections:
                self._write_sheet(writer, "sections", pd.DataFrame(self._section_rows(sections)))
            if indicators and not records:
                self._write_sheet(writer, "indicators", pd.DataFrame(self._indicator_rows(indicators)))

            if datasets:
                self._write_sheet(writer, "dataset_index", pd.DataFrame(self._dataset_index_rows(datasets)))
                for index, dataset in enumerate(datasets, start=1):
                    dataframe = pd.DataFrame(dataset.get("records", []))
                    if dataframe.empty:
                        continue
                    dataset_name = str(dataset.get("name") or f"dataset_{index}")
                    self._write_sheet(writer, f"table_{index:02d}_{dataset_name}", dataframe)

            self._format_workbook(writer)

        return self._build_result("excel", output_path)

    def _generate_word(self, payload: dict[str, Any]) -> dict[str, Any]:
        if Document is None:
            raise RuntimeError("python-docx 未安装，无法生成 Word 文件")

        output_path = self._build_output_path(payload, "docx")
        natural_language = payload.get("natural_language", {})
        structured_data = payload.get("structured_data", {})
        title = (
            natural_language.get("title")
            or payload.get("document_summary", {}).get("title")
            or payload.get("canonical_data", {}).get("document_meta", {}).get("title")
            or "统计分析报告"
        )

        document = Document()
        document.add_heading(title, level=1)
        document.add_paragraph(f"源文件: {payload.get('source_file', 'unknown')}")
        document.add_paragraph(f"生成时间: {payload.get('generated_at', '')}")
        document.add_paragraph(f"转换模式: {payload.get('operation_mode', 'natural_to_structured')}")
        document.add_paragraph(f"文档类型: {self._document_type_label(structured_data.get('document_type'))}")

        content = str(natural_language.get("content") or "").strip()
        if content:
            document.add_heading("统计分析报告正文", level=2)
            for paragraph in content.split("\n\n"):
                paragraph = paragraph.strip()
                if paragraph:
                    document.add_paragraph(paragraph)
        else:
            document.add_heading("结构化摘要", level=2)
            for paragraph in self._fallback_report_paragraphs(payload):
                document.add_paragraph(paragraph)

        chart_outputs = payload.get("chart_outputs", []) if isinstance(payload.get("chart_outputs"), list) else []
        chart_render_specs = payload.get("chart_render_specs", []) if isinstance(payload.get("chart_render_specs"), list) else []
        if chart_render_specs or chart_outputs:
            document.add_heading("图表", level=2)
            charts_to_embed = chart_render_specs or chart_outputs
            for index, item in enumerate(charts_to_embed):
                title_text = str(item.get("title") or item.get("id") or "图表")
                document.add_paragraph(title_text)
                embedded = False
                if Inches is not None and chart_render_specs and index < len(chart_render_specs):
                    image_buffer = self.chart_renderer.render_buffer(chart_render_specs[index])
                    if image_buffer is not None:
                        try:
                            document.add_picture(image_buffer, width=Inches(6.2))
                            embedded = True
                        finally:
                            image_buffer.close()
                if not embedded:
                    chart_path = str(item.get("path") or "")
                    if Inches is not None and chart_path and os.path.exists(chart_path):
                        try:
                            document.add_picture(chart_path, width=Inches(6.2))
                            embedded = True
                        except Exception:
                            embedded = False
                if not embedded:
                    chart_path = str(item.get("path") or "")
                    if chart_path:
                        document.add_paragraph(f"图表路径: {chart_path}")
                reason = str(item.get("reason") or "").strip()
                if reason:
                    document.add_paragraph(f"图表说明: {reason}")

        document.save(output_path)
        return self._build_result("word", output_path)

    def _normalize_payload(self, data: Any) -> dict[str, Any]:
        pipeline = data if isinstance(data, dict) else {}
        parse_result = pipeline.get("parse", {}) if isinstance(pipeline.get("parse", {}), dict) else {}
        semantic_result = pipeline.get("semantic", {}) if isinstance(pipeline.get("semantic", {}), dict) else {}
        validation_result = pipeline.get("validation", {}) if isinstance(pipeline.get("validation", {}), dict) else {}
        structured_data = semantic_result.get("structured_data", semantic_result) if isinstance(semantic_result, dict) else {}
        canonical_data = semantic_result.get("canonical_data", pipeline.get("canonical_data", {})) if isinstance(semantic_result, dict) else {}
        natural_language = pipeline.get("natural_language", {}) if isinstance(pipeline.get("natural_language", {}), dict) else {}
        document_summary = semantic_result.get("document_summary") if isinstance(semantic_result.get("document_summary"), dict) else {}
        if not document_summary and isinstance(structured_data, dict):
            document_summary = structured_data.get("document_summary", {}) if isinstance(structured_data.get("document_summary"), dict) else {}
        if not document_summary and isinstance(canonical_data, dict):
            document_summary = canonical_data.get("document_summary", {}) if isinstance(canonical_data.get("document_summary"), dict) else {}
        detail_groups = semantic_result.get("detail_groups") if isinstance(semantic_result.get("detail_groups"), dict) else {}
        if not detail_groups and isinstance(structured_data, dict):
            detail_groups = structured_data.get("detail_groups", {}) if isinstance(structured_data.get("detail_groups"), dict) else {}
        if not detail_groups and isinstance(canonical_data, dict):
            detail_groups = canonical_data.get("detail_groups", {}) if isinstance(canonical_data.get("detail_groups"), dict) else {}

        return {
            "source_file": pipeline.get("source_file") or parse_result.get("metadata", {}).get("file_name") or "output",
            "generated_at": datetime.now().isoformat(),
            "operation_mode": pipeline.get("operation_mode", "natural_to_structured"),
            "source_type": parse_result.get("type", structured_data.get("source_type", "unknown")) if isinstance(structured_data, dict) else parse_result.get("type", "unknown"),
            "parse_summary": {
                "elements": parse_result.get("elements", 0),
                "table_count": len(parse_result.get("tables", [])) if isinstance(parse_result.get("tables", []), list) else 0,
                "block_count": len(parse_result.get("blocks", [])) if isinstance(parse_result.get("blocks", []), list) else 0,
            },
            "document_understanding": pipeline.get("understanding", {}),
            "structured_data": structured_data if isinstance(structured_data, dict) else {},
            "canonical_data": canonical_data if isinstance(canonical_data, dict) else {},
            "document_summary": document_summary,
            "detail_groups": detail_groups,
            "preprocessing": semantic_result.get("preprocessing", structured_data.get("preprocessing", {})) if isinstance(structured_data, dict) else {},
            "routing": semantic_result.get("routing", structured_data.get("routing", {})) if isinstance(structured_data, dict) else {},
            "prompting": semantic_result.get("prompting", structured_data.get("prompting", {})) if isinstance(structured_data, dict) else {},
            "postprocessing": semantic_result.get("postprocessing", structured_data.get("postprocessing", {})) if isinstance(structured_data, dict) else {},
            "models": semantic_result.get("models", {}) if isinstance(semantic_result.get("models"), dict) else {},
            "natural_language": natural_language,
            "report_facts": natural_language.get("report_facts", {}) if isinstance(natural_language.get("report_facts"), dict) else {},
            "report_plan": natural_language.get("report_plan", {}) if isinstance(natural_language.get("report_plan"), dict) else {},
            "chart_specs": natural_language.get("chart_specs", []) if isinstance(natural_language.get("chart_specs"), list) else [],
            "chart_render_specs": natural_language.get("chart_render_specs", []) if isinstance(natural_language.get("chart_render_specs"), list) else [],
            "chart_outputs": natural_language.get("chart_outputs", []) if isinstance(natural_language.get("chart_outputs"), list) else [],
            "validation": validation_result,
        }

    def _overview_rows(self, payload: dict[str, Any]) -> list[dict[str, Any]]:
        summary = payload.get("document_summary", {}) if isinstance(payload.get("document_summary"), dict) else {}
        natural_language = payload.get("natural_language", {}) if isinstance(payload.get("natural_language"), dict) else {}
        report_facts = payload.get("report_facts", {}) if isinstance(payload.get("report_facts"), dict) else {}
        forbidden = natural_language.get("forbidden_expression_check", {}) if isinstance(natural_language.get("forbidden_expression_check"), dict) else {}
        structured_data = payload.get("structured_data", {}) if isinstance(payload.get("structured_data"), dict) else {}
        canonical_meta = payload.get("canonical_data", {}).get("document_meta", {}) if isinstance(payload.get("canonical_data", {}).get("document_meta", {}), dict) else {}
        records = structured_data.get("records", []) if isinstance(structured_data.get("records"), list) else []
        indicators = structured_data.get("indicators", []) if isinstance(structured_data.get("indicators"), list) else []
        detail_groups = payload.get("detail_groups", {}).get("groups", []) if isinstance(payload.get("detail_groups", {}).get("groups", []), list) else []

        rows = [
            {"分组": "概览", "字段": "源文件", "值": payload.get("source_file", "")},
            {"分组": "概览", "字段": "生成时间", "值": payload.get("generated_at", "")},
            {"分组": "概览", "字段": "标题", "值": natural_language.get("title") or summary.get("title") or canonical_meta.get("title") or structured_data.get("title", "")},
            {"分组": "概览", "字段": "文档类型", "值": self._document_type_label(summary.get("document_type") or structured_data.get("document_type"))},
            {"分组": "概览", "字段": "转换模式", "值": payload.get("operation_mode", "")},
            {"分组": "概览", "字段": "领域", "值": report_facts.get("domain", self._domain_text(summary.get("domain_tags", [])))},
            {"分组": "概览", "字段": "期间", "值": report_facts.get("period", summary.get("year", ""))},
            {"分组": "概览", "字段": "地区", "值": summary.get("region") or canonical_meta.get("region") or structured_data.get("region", "")},
            {"分组": "概览", "字段": "总体摘要", "值": summary.get("overall_summary", "")},
            {"分组": "概览", "字段": "专题数", "值": len(report_facts.get("topics", [])) if isinstance(report_facts.get("topics"), list) else 0},
            {"分组": "概览", "字段": "图表数", "值": len(payload.get("chart_outputs", [])) if isinstance(payload.get("chart_outputs", []), list) else 0},
            {"分组": "概览", "字段": "记录数", "值": len(records)},
            {"分组": "概览", "字段": "指标数", "值": len(indicators)},
            {"分组": "概览", "字段": "详细子表数", "值": len(detail_groups)},
            {"分组": "概览", "字段": "校验通过", "值": payload.get("validation", {}).get("is_valid", False)},
            {"分组": "概览", "字段": "正文禁用词命中", "值": ", ".join(str(item) for item in forbidden.get("matched", []))},
        ]
        for index, item in enumerate(report_facts.get("highlights", []) or [], start=1):
            rows.append({"分组": "亮点", "字段": f"亮点{index}", "值": item})
        for index, item in enumerate(report_facts.get("risks", []) or [], start=1):
            rows.append({"分组": "风险", "字段": f"风险{index}", "值": item})
        return [row for row in rows if row.get("值") not in {None, ""} or isinstance(row.get("值"), bool) or row.get("字段") in {"专题数", "图表数", "记录数", "指标数", "详细子表数"}]

    def _detail_index_rows(self, payload: dict[str, Any]) -> list[dict[str, Any]]:
        detail_groups = payload.get("detail_groups", {}) if isinstance(payload.get("detail_groups"), dict) else {}
        groups = detail_groups.get("groups", []) if isinstance(detail_groups.get("groups"), list) else []
        rows = []
        for index, group in enumerate(groups, start=1):
            rows.append(
                {
                    "序号": index,
                    "sheet_name": group.get("sheet_name", ""),
                    "分类": group.get("sheet_category", ""),
                    "分组值": group.get("group_value", ""),
                    "来源类型": group.get("source_kind", ""),
                    "记录数": group.get("record_count", 0),
                    "说明": group.get("description", ""),
                }
            )
        return rows

    def _pipeline_plan_rows(self, payload: dict[str, Any]) -> list[dict[str, Any]]:
        models = payload.get("models", {}) if isinstance(payload.get("models"), dict) else {}
        routing = payload.get("routing", {}) if isinstance(payload.get("routing"), dict) else {}
        preprocessing = payload.get("preprocessing", {}) if isinstance(payload.get("preprocessing"), dict) else {}
        prompting = payload.get("prompting", {}) if isinstance(payload.get("prompting"), dict) else {}
        report_prompting = payload.get("natural_language", {}).get("prompting", {}) if isinstance(payload.get("natural_language", {}).get("prompting", {}), dict) else {}
        postprocessing = payload.get("postprocessing", {}) if isinstance(payload.get("postprocessing"), dict) else {}
        rows = []

        for key, value in models.items():
            if value not in (None, "", []):
                rows.append({"分组": "模型", "字段": key, "值": value})
        for key, value in routing.items():
            if value not in (None, "", []):
                rows.append({"分组": "路由", "字段": key, "值": self._stringify_value(value)})

        stats = preprocessing.get("stats", {}) if isinstance(preprocessing.get("stats"), dict) else {}
        signals = preprocessing.get("signals", {}) if isinstance(preprocessing.get("signals"), dict) else {}
        for key, value in stats.items():
            if value not in (None, "", []):
                rows.append({"分组": "预处理", "字段": key, "值": value})
        for key, value in signals.items():
            if value not in (None, "", []):
                rows.append({"分组": "预处理", "字段": key, "值": self._stringify_value(value)})

        for stage, node in prompting.items():
            if not isinstance(node, dict):
                continue
            for key, value in node.items():
                if value not in (None, "", []):
                    rows.append({"分组": f"语义提示:{stage}", "字段": key, "值": self._stringify_value(value)})
        for stage, node in report_prompting.items():
            if not isinstance(node, dict):
                continue
            for key, value in node.items():
                if value not in (None, "", []):
                    rows.append({"分组": f"报告提示:{stage}", "字段": key, "值": self._stringify_value(value)})

        for key, value in postprocessing.items():
            if value not in (None, "", []):
                rows.append({"分组": "后处理", "字段": key, "值": self._stringify_value(value)})
        return rows

    def _report_fact_rows(self, report_facts: dict[str, Any]) -> list[dict[str, Any]]:
        if not isinstance(report_facts, dict):
            return []
        rows = [
            {"分组": "报告事实", "字段": "标题", "值": report_facts.get("title", "")},
            {"分组": "报告事实", "字段": "期间", "值": report_facts.get("period", "")},
            {"分组": "报告事实", "字段": "领域", "值": report_facts.get("domain", "")},
            {"分组": "报告事实", "字段": "专题数", "值": len(report_facts.get("topics", [])) if isinstance(report_facts.get("topics"), list) else 0},
        ]
        for index, item in enumerate(report_facts.get("highlights", []) or [], start=1):
            rows.append({"分组": "亮点", "字段": f"亮点{index}", "值": item})
        for index, item in enumerate(report_facts.get("risks", []) or [], start=1):
            rows.append({"分组": "风险", "字段": f"风险{index}", "值": item})
        return [row for row in rows if row.get("值") not in {None, ""} or row.get("字段") == "专题数"]

    def _report_topic_rows(self, report_facts: dict[str, Any]) -> list[dict[str, Any]]:
        topics = report_facts.get("topics", []) if isinstance(report_facts.get("topics"), list) else []
        rows = []
        for index, topic in enumerate(topics, start=1):
            metrics = topic.get("summary_metric", []) if isinstance(topic.get("summary_metric"), list) else []
            breakdowns = topic.get("breakdowns", []) if isinstance(topic.get("breakdowns"), list) else []
            time_series = topic.get("time_series", []) if isinstance(topic.get("time_series"), list) else []
            anomalies = topic.get("anomalies", []) if isinstance(topic.get("anomalies"), list) else []
            insights = topic.get("insights", []) if isinstance(topic.get("insights"), list) else []
            metric_rows = metrics[:4] if metrics else [{}]
            for metric in metric_rows:
                rows.append(
                    {
                        "专题序号": index,
                        "专题": topic.get("name", ""),
                        "指标": metric.get("indicator", ""),
                        "数值": metric.get("value"),
                        "单位": metric.get("unit", ""),
                        "期间": metric.get("period", ""),
                        "变化值": metric.get("change_value"),
                        "变化说明": metric.get("change_text", ""),
                        "拆分维度": breakdowns[0].get("dimension", "") if breakdowns else "",
                        "时间序列点数": len(time_series[0].get("points", [])) if time_series else 0,
                        "异常提示": anomalies[0].get("description", "") if anomalies else "",
                        "研判": insights[0] if insights else "",
                    }
                )
        return rows

    def _report_plan_rows(self, report_plan: dict[str, Any]) -> list[dict[str, Any]]:
        sections = report_plan.get("sections", []) if isinstance(report_plan.get("sections"), list) else []
        rows = []
        for index, section in enumerate(sections, start=1):
            focus_metrics = section.get("focus_metrics", []) if isinstance(section.get("focus_metrics"), list) else []
            rows.append(
                {
                    "序号": index,
                    "章节ID": section.get("id", ""),
                    "章节标题": section.get("title", ""),
                    "写作目的": section.get("purpose", ""),
                    "专题引用": " | ".join(str(item) for item in section.get("topic_refs", []) or []),
                    "重点指标": " | ".join(str(metric.get("indicator") or metric.get("name") or "") for metric in focus_metrics[:6]),
                    "图表数量": len(section.get("chart_specs", []) or []),
                    "图表标题": " | ".join(str(chart.get("title", "")) for chart in section.get("chart_specs", []) or []),
                }
            )
        return rows

    def _chart_spec_rows(self, chart_specs: list[dict[str, Any]]) -> list[dict[str, Any]]:
        rows = []
        for index, spec in enumerate(chart_specs if isinstance(chart_specs, list) else [], start=1):
            if not isinstance(spec, dict):
                continue
            series = spec.get("series", []) if isinstance(spec.get("series"), list) else []
            point_count = 0
            for item in series:
                if isinstance(item, dict):
                    point_count += len(item.get("data", []) if isinstance(item.get("data"), list) else [])
            rows.append(
                {
                    "序号": index,
                    "图表类型": spec.get("type", "none"),
                    "标题": spec.get("title", ""),
                    "X字段": spec.get("x_field", ""),
                    "Y字段": spec.get("y_field", ""),
                    "系列数": len(series),
                    "数据点数": point_count,
                    "生成原因": spec.get("reason", ""),
                }
            )
        return rows

    def _chart_output_rows(self, chart_outputs: list[dict[str, Any]]) -> list[dict[str, Any]]:
        rows = []
        for index, item in enumerate(chart_outputs if isinstance(chart_outputs, list) else [], start=1):
            if not isinstance(item, dict):
                continue
            rows.append(
                {
                    "序号": index,
                    "图表ID": item.get("id", ""),
                    "图表类型": item.get("type", ""),
                    "标题": item.get("title", ""),
                    "文件路径": item.get("path", ""),
                    "文件大小": item.get("size", ""),
                    "生成原因": item.get("reason", ""),
                    "错误": item.get("error", ""),
                }
            )
        return rows

    def _record_rows(
        self,
        records: list[dict[str, Any]],
        structured_data: dict[str, Any],
        canonical_data: dict[str, Any],
    ) -> list[dict[str, Any]]:
        meta = canonical_data.get("document_meta", {}) if isinstance(canonical_data.get("document_meta"), dict) else {}
        rows = []
        for record in records:
            rows.append(
                {
                    "记录ID": record.get("record_id", ""),
                    "记录类型": record.get("record_type", ""),
                    "文档标题": record.get("document_title") or meta.get("title") or structured_data.get("title", ""),
                    "时间": record.get("time_reference", ""),
                    "地区": record.get("location_reference", ""),
                    "章节": record.get("section", ""),
                    "指标名称": record.get("indicator_name", ""),
                    "指标类型": record.get("metric_type", ""),
                    "数值": record.get("value"),
                    "单位": record.get("unit", ""),
                    "同比(%)": record.get("yoy_percent"),
                    "说明": record.get("comparison_text", ""),
                    "值文本": record.get("value_text", ""),
                    "来源句子": record.get("source_sentence", ""),
                }
            )
        return rows

    def _section_rows(self, sections: list[dict[str, Any]]) -> list[dict[str, Any]]:
        rows = []
        for index, section in enumerate(sections, start=1):
            rows.append(
                {
                    "序号": index,
                    "章节": section.get("title", f"section_{index}"),
                    "摘要": str(section.get("content") or section.get("summary") or "")[:240],
                    "原文字数": len(str(section.get("content") or section.get("summary") or "")),
                }
            )
        return rows

    def _indicator_rows(self, indicators: list[dict[str, Any]]) -> list[dict[str, Any]]:
        rows = []
        for index, indicator in enumerate(indicators, start=1):
            rows.append(
                {
                    "序号": index,
                    "章节": indicator.get("section_title", ""),
                    "指标名称": indicator.get("name", ""),
                    "数值": indicator.get("value"),
                    "单位": indicator.get("unit", ""),
                    "同比变化(%)": indicator.get("yoy"),
                    "说明": indicator.get("comparison_text", ""),
                    "来源句子": indicator.get("statement", ""),
                }
            )
        return rows

    def _dataset_index_rows(self, datasets: list[dict[str, Any]]) -> list[dict[str, Any]]:
        rows = []
        for index, dataset in enumerate(datasets, start=1):
            headers = dataset.get("headers", []) if isinstance(dataset.get("headers"), list) else []
            rows.append(
                {
                    "序号": index,
                    "数据表": dataset.get("name", f"dataset_{index}"),
                    "记录数": dataset.get("row_count", len(dataset.get("records", []) if isinstance(dataset.get("records"), list) else [])),
                    "字段数": dataset.get("column_count", len(headers)),
                    "字段列表": "、".join(str(item) for item in headers[:20]),
                }
            )
        return rows

    def _metadata_rows(self, structured_data: dict[str, Any], canonical_data: dict[str, Any]) -> list[dict[str, Any]]:
        meta = canonical_data.get("document_meta", {}) if isinstance(canonical_data.get("document_meta"), dict) else {}
        rows = [
            {"分组": "文档信息", "字段": "标题", "值": meta.get("title") or structured_data.get("title", "")},
            {"分组": "文档信息", "字段": "文档类型", "值": self._document_type_label(meta.get("document_type") or structured_data.get("document_type"))},
            {"分组": "文档信息", "字段": "地区", "值": meta.get("region") or structured_data.get("region", "")},
            {"分组": "文档信息", "字段": "年份", "值": meta.get("year") or structured_data.get("year", "")},
            {"分组": "文档信息", "字段": "发布日期", "值": meta.get("publish_date") or structured_data.get("publish_date", "")},
            {"分组": "文档信息", "字段": "行业领域", "值": self._domain_text(meta.get("domain_tags", structured_data.get("domain_tags", [])))},
        ]
        key_values = structured_data.get("key_values", {}) if isinstance(structured_data.get("key_values"), dict) else {}
        for key, value in list(key_values.items())[:30]:
            rows.append({"分组": "原文字段", "字段": str(key), "值": self._stringify_value(value)})
        return [row for row in rows if row.get("值") not in {None, ""}]

    def _fallback_report_paragraphs(self, payload: dict[str, Any]) -> list[str]:
        report_facts = payload.get("report_facts", {}) if isinstance(payload.get("report_facts"), dict) else {}
        paragraphs = []
        for topic in report_facts.get("topics", [])[:4] if isinstance(report_facts.get("topics", []), list) else []:
            topic_name = str(topic.get("name") or "重点指标")
            metrics = topic.get("summary_metric", []) if isinstance(topic.get("summary_metric"), list) else []
            if not metrics:
                continue
            metric = metrics[0]
            value = self._stringify_value(metric.get("value"))
            unit = str(metric.get("unit") or "")
            change_value = metric.get("change_value")
            if change_value not in {None, ""}:
                paragraphs.append(f"{topic_name}为{value}{unit}，变化值为{change_value}。")
            else:
                paragraphs.append(f"{topic_name}为{value}{unit}。")
        if not paragraphs:
            paragraphs.append("当前结果已生成结构化事实，但未形成可写入的报告正文。")
        return paragraphs

    def _fallback_report_text(self, payload: dict[str, Any]) -> str:
        title = self._base_name(payload)
        paragraphs = self._fallback_report_paragraphs(payload)
        return title + "\n\n" + "\n\n".join(paragraphs)

    def _pdf_source_text(self, payload: dict[str, Any]) -> str:
        title = self._base_name(payload)
        content = str(payload.get("natural_language", {}).get("content") or "").strip()
        if not content:
            content = self._fallback_report_text(payload)
        header_lines = [title, "", f"源文件: {payload.get('source_file', '')}", f"转换模式: {payload.get('operation_mode', '')}"]
        document_type = payload.get("structured_data", {}).get("document_type", "")
        if document_type:
            header_lines.append(f"文档类型: {self._document_type_label(document_type)}")
        return "\n".join(header_lines) + "\n\n" + content

    def _pdf_lines(self, text: str) -> list[str]:
        normalized = str(text or "").replace("\r\n", "\n").replace("\r", "\n")
        paragraphs = normalized.split("\n\n")
        lines: list[str] = []
        for paragraph in paragraphs:
            raw_lines = paragraph.split("\n")
            for raw_line in raw_lines:
                stripped = raw_line.strip()
                if not stripped:
                    lines.append("")
                    continue
                lines.extend(self._wrap_pdf_text(stripped))
            lines.append("")
        while lines and not lines[-1]:
            lines.pop()
        return lines or [""]

    def _wrap_pdf_text(self, text: str, max_units: int = 52) -> list[str]:
        wrapped: list[str] = []
        current = ""
        units = 0
        for char in text.replace("\t", "    "):
            width = self._char_display_width(char)
            if current and units + width > max_units:
                wrapped.append(current)
                current = char
                units = width
            else:
                current += char
                units += width
        if current:
            wrapped.append(current)
        return wrapped or [""]

    def _char_display_width(self, char: str) -> int:
        return 1 if unicodedata.east_asian_width(char) in {"Na", "H", "N"} else 2

    def _build_pdf_document(self, lines: list[str]) -> bytes:
        page_width = 595
        page_height = 842
        left_margin = 50
        top_margin = 800
        line_height = 16
        lines_per_page = 45

        pages: list[list[str]] = []
        current_page: list[str] = []
        for line in lines:
            current_page.append(line)
            if len(current_page) >= lines_per_page:
                pages.append(current_page)
                current_page = []
        if current_page or not pages:
            pages.append(current_page)

        objects: dict[int, bytes] = {}
        catalog_num = 1
        pages_num = 2
        font_num = 3
        cidfont_num = 4
        next_num = 5

        objects[cidfont_num] = (
            "<< /Type /Font /Subtype /CIDFontType0 /BaseFont /STSong-Light "
            "/CIDSystemInfo << /Registry (Adobe) /Ordering (GB1) /Supplement 4 >> /DW 1000 >>"
        ).encode("latin-1")
        objects[font_num] = (
            f"<< /Type /Font /Subtype /Type0 /BaseFont /STSong-Light /Encoding /UniGB-UCS2-H /DescendantFonts [{cidfont_num} 0 R] >>"
        ).encode("latin-1")

        page_nums: list[int] = []
        for page_lines in pages:
            stream = self._build_pdf_page_stream(page_lines, left_margin, top_margin, line_height)
            content_num = next_num
            page_num = next_num + 1
            next_num += 2
            objects[content_num] = b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream"
            objects[page_num] = (
                f"<< /Type /Page /Parent {pages_num} 0 R /MediaBox [0 0 {page_width} {page_height}] "
                f"/Resources << /Font << /F1 {font_num} 0 R >> >> /Contents {content_num} 0 R >>"
            ).encode("latin-1")
            page_nums.append(page_num)

        kids = " ".join(f"{num} 0 R" for num in page_nums)
        objects[pages_num] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_nums)} >>".encode("latin-1")
        objects[catalog_num] = f"<< /Type /Catalog /Pages {pages_num} 0 R >>".encode("latin-1")

        buffer = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
        offsets = {0: 0}
        for obj_num in sorted(objects):
            offsets[obj_num] = len(buffer)
            buffer.extend(f"{obj_num} 0 obj\n".encode("latin-1"))
            buffer.extend(objects[obj_num])
            buffer.extend(b"\nendobj\n")

        xref_offset = len(buffer)
        buffer.extend(f"xref\n0 {max(objects) + 1}\n".encode("latin-1"))
        buffer.extend(b"0000000000 65535 f \n")
        for obj_num in range(1, max(objects) + 1):
            offset = offsets.get(obj_num, 0)
            buffer.extend(f"{offset:010d} 00000 n \n".encode("latin-1"))
        buffer.extend(
            (
                f"trailer\n<< /Size {max(objects) + 1} /Root {catalog_num} 0 R >>\n"
                f"startxref\n{xref_offset}\n%%EOF"
            ).encode("latin-1")
        )
        return bytes(buffer)

    def _build_pdf_page_stream(self, lines: list[str], left_margin: int, top_margin: int, line_height: int) -> bytes:
        commands = ["BT", "/F1 11 Tf", f"1 0 0 1 {left_margin} {top_margin} Tm", f"{line_height} TL"]
        first_line = True
        for line in lines:
            if first_line:
                first_line = False
            else:
                commands.append("T*")
            if line:
                commands.append(f"<{self._pdf_hex_text(line)}> Tj")
        commands.append("ET")
        return "\n".join(commands).encode("latin-1")

    def _pdf_hex_text(self, text: str) -> str:
        encoded = text.encode("utf-16-be")
        return "FEFF" + encoded.hex().upper()

    def _write_sheet(self, writer: pd.ExcelWriter, sheet_name: str, dataframe: pd.DataFrame) -> None:
        if dataframe is None or dataframe.empty:
            return
        safe_name = self._safe_sheet_name(sheet_name, set(writer.book.sheetnames))
        dataframe.to_excel(writer, sheet_name=safe_name, index=False)

    def _format_workbook(self, writer: pd.ExcelWriter) -> None:
        if Font is None or PatternFill is None or get_column_letter is None:
            return
        workbook = writer.book
        for sheet in workbook.worksheets:
            if sheet.max_row <= 0 or sheet.max_column <= 0:
                continue
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill(fill_type="solid", fgColor=HEADER_FILL)
            for column_index in range(1, sheet.max_column + 1):
                values = [self._stringify_value(sheet.cell(row=row_index, column=column_index).value) for row_index in range(1, sheet.max_row + 1)]
                max_length = min(max((len(value) for value in values), default=8) + 2, 36)
                sheet.column_dimensions[get_column_letter(column_index)].width = max(10, max_length)

    def _build_output_path(self, payload: dict[str, Any], extension: str) -> str:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        base_name = self._safe_file_name(self._base_name(payload))
        return os.path.join(OUTPUT_DIR, f"{base_name}_{timestamp}.{extension}")

    def _build_result(self, output_format: str, output_path: str) -> dict[str, Any]:
        return {
            "format": output_format,
            "path": output_path,
            "size": os.path.getsize(output_path) if os.path.exists(output_path) else 0,
        }

    def _base_name(self, payload: dict[str, Any]) -> str:
        natural_language = payload.get("natural_language", {}) if isinstance(payload.get("natural_language"), dict) else {}
        document_summary = payload.get("document_summary", {}) if isinstance(payload.get("document_summary"), dict) else {}
        canonical_meta = payload.get("canonical_data", {}).get("document_meta", {}) if isinstance(payload.get("canonical_data", {}).get("document_meta", {}), dict) else {}
        source_file = str(payload.get("source_file") or "output")
        source_name = os.path.splitext(os.path.basename(source_file))[0]
        return str(
            natural_language.get("title")
            or document_summary.get("title")
            or canonical_meta.get("title")
            or payload.get("structured_data", {}).get("title")
            or source_name
            or "output"
        )

    def _safe_sheet_name(self, name: str, existing: set[str]) -> str:
        cleaned = re.sub(r"[\\/*?:\[\]]+", "_", str(name or "sheet")).strip() or "sheet"
        cleaned = cleaned[:31]
        if cleaned not in existing:
            return cleaned
        suffix = 1
        while True:
            candidate = f"{cleaned[:28]}_{suffix}"
            if candidate not in existing:
                return candidate
            suffix += 1

    def _safe_file_name(self, name: str) -> str:
        cleaned = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff._-]+", "_", str(name or "output")).strip("._")
        return cleaned[:80] or "output"

    def _document_type_label(self, document_type: Any) -> str:
        return DOCUMENT_TYPE_LABELS.get(str(document_type or "unknown"), str(document_type or "unknown"))

    def _domain_text(self, domains: Any) -> str:
        if isinstance(domains, list):
            return "、".join(DOMAIN_LABELS.get(str(item), str(item)) for item in domains if item)
        if domains:
            return DOMAIN_LABELS.get(str(domains), str(domains))
        return ""

    def _stringify_value(self, value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, (dict, list)):
            return json.dumps(value, ensure_ascii=False)
        return str(value)
