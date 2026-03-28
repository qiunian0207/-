from __future__ import annotations

import json
from pathlib import Path
from typing import Any

try:
    from config.config import TABLE_PREVIEW_ROWS, TEXT_PREVIEW_CHARS
except ImportError:
    from doc_processing.config.config import TABLE_PREVIEW_ROWS, TEXT_PREVIEW_CHARS

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None


IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".tiff"}
TEXT_EXTENSIONS = {".txt", ".md"}
JSON_EXTENSIONS = {".json"}


class DocumentParser:
    def parse(self, file_path: str) -> dict[str, Any]:
        path = Path(file_path)
        extension = path.suffix.lower()

        try:
            if extension == ".xlsx":
                return self._parse_xlsx(path)
            if extension == ".docx":
                return self._parse_docx(path)
            if extension == ".pdf":
                return self._parse_pdf(path)
            if extension in IMAGE_EXTENSIONS:
                return self._parse_image(path)
            if extension in TEXT_EXTENSIONS:
                return self._parse_text_file(path)
            if extension in JSON_EXTENSIONS:
                return self._parse_json_file(path)
            raise ValueError(f"Unsupported file format: {extension}")
        except Exception as exc:
            return {
                "error": str(exc),
                "type": extension.lstrip(".") or "unknown",
                "metadata": self._base_metadata(path),
            }

    def _parse_text_file(self, path: Path) -> dict[str, Any]:
        text = path.read_text(encoding="utf-8", errors="ignore")
        paragraphs = [line.strip() for line in text.splitlines() if line.strip()]
        blocks = [
            {
                "id": f"line_{index}",
                "type": "paragraph",
                "text": paragraph,
                "source_anchor": {"line_index": index},
            }
            for index, paragraph in enumerate(paragraphs, start=1)
        ]
        return {
            "type": path.suffix.lstrip("."),
            "content": text,
            "elements": len(paragraphs),
            "tables": [],
            "blocks": blocks,
            "json_data": None,
            "metadata": {
                **self._base_metadata(path),
                "title": paragraphs[0] if paragraphs else path.stem,
                "text_length": len(text),
            },
        }

    def _parse_json_file(self, path: Path) -> dict[str, Any]:
        raw_text = path.read_text(encoding="utf-8", errors="ignore")
        payload = json.loads(raw_text)
        tables: list[dict[str, Any]] = []
        blocks: list[dict[str, Any]] = []
        title = path.stem

        if isinstance(payload, list) and payload and all(isinstance(item, dict) for item in payload):
            headers = []
            for record in payload:
                for key in record.keys():
                    if key not in headers:
                        headers.append(str(key))
            rows = [headers]
            for record in payload:
                rows.append([self._clean_text(record.get(header)) for header in headers])
            table_payload = self._build_table_payload("records", rows, {"source_file": path.name})
            if table_payload:
                tables.append(table_payload)
        elif isinstance(payload, dict):
            if isinstance(payload.get("document_meta"), dict):
                title = str(payload["document_meta"].get("title") or title)
            for index, (key, value) in enumerate(payload.items(), start=1):
                blocks.append(
                    {
                        "id": f"json_field_{index}",
                        "type": "json_field",
                        "text": f"{key}: {value}",
                        "source_anchor": {"field": key},
                    }
                )

        if not blocks and raw_text.strip():
            blocks.append(
                {
                    "id": "json_text",
                    "type": "json",
                    "text": raw_text[:TEXT_PREVIEW_CHARS],
                    "source_anchor": {"file_name": path.name},
                }
            )

        return {
            "type": "json",
            "content": raw_text,
            "elements": len(blocks) or (len(payload) if hasattr(payload, "__len__") else 1),
            "tables": tables,
            "blocks": blocks,
            "json_data": payload,
            "metadata": {
                **self._base_metadata(path),
                "title": title,
                "text_length": len(raw_text),
                "record_count": len(payload) if isinstance(payload, list) else 0,
            },
        }

    def _parse_docx(self, path: Path) -> dict[str, Any]:
        warnings: list[str] = []
        paragraphs: list[str] = []
        tables: list[dict[str, Any]] = []
        blocks: list[dict[str, Any]] = []

        if DocxDocument is not None:
            document = DocxDocument(str(path))
            for index, paragraph in enumerate(document.paragraphs, start=1):
                text = paragraph.text.strip() if paragraph.text else ""
                if not text:
                    continue
                paragraphs.append(text)
                style_name = getattr(getattr(paragraph, "style", None), "name", "") or ""
                blocks.append(
                    {
                        "id": f"paragraph_{index}",
                        "type": "paragraph",
                        "text": text,
                        "style": style_name,
                        "source_anchor": {"paragraph_index": index, "style": style_name},
                    }
                )
            for index, table in enumerate(document.tables, start=1):
                rows = []
                for row in table.rows:
                    rows.append([self._clean_text(cell.text) for cell in row.cells])
                table_payload = self._build_table_payload(
                    f"table_{index}",
                    rows,
                    {"table_index": index},
                )
                if table_payload:
                    tables.append(table_payload)
        else:
            warnings.append("python-docx 未安装，DOCX 仅能尝试通过 unstructured 降级读取。")
            extracted = self._partition_docx(path)
            if extracted:
                paragraphs = extracted
                blocks = [
                    {
                        "id": f"paragraph_{index}",
                        "type": "paragraph",
                        "text": paragraph,
                        "source_anchor": {"paragraph_index": index},
                    }
                    for index, paragraph in enumerate(extracted, start=1)
                ]

        content = "\n".join(paragraphs).strip()
        if not content and tables:
            content = self._tables_to_preview_text(tables)

        return {
            "type": "docx",
            "content": content,
            "elements": len(paragraphs) + sum(table["row_count"] for table in tables),
            "tables": tables,
            "blocks": blocks,
            "json_data": None,
            "metadata": {
                **self._base_metadata(path),
                "title": paragraphs[0] if paragraphs else path.stem,
                "table_count": len(tables),
                "text_length": len(content),
            },
            "warnings": warnings,
        }

    def _parse_xlsx(self, path: Path) -> dict[str, Any]:
        if openpyxl is None:
            raise RuntimeError("openpyxl 未安装，无法解析 XLSX 文件")

        workbook = openpyxl.load_workbook(path, data_only=True)
        tables: list[dict[str, Any]] = []
        blocks: list[dict[str, Any]] = []

        for worksheet in workbook.worksheets:
            rows = []
            for row in worksheet.iter_rows(values_only=True):
                rows.append([self._clean_text(value) for value in row])
            table_payload = self._build_table_payload(worksheet.title, rows, {"sheet": worksheet.title})
            if table_payload:
                tables.append(table_payload)
                blocks.append(
                    {
                        "id": f"sheet_{len(blocks) + 1}",
                        "type": "table",
                        "text": f"{worksheet.title}: {table_payload['row_count']} 行, {table_payload['column_count']} 列",
                        "source_anchor": {"sheet": worksheet.title},
                    }
                )

        total_records = sum(table["row_count"] for table in tables)
        schema = sorted({header for table in tables for header in table["headers"]})

        return {
            "type": "xlsx",
            "content": self._tables_to_preview_text(tables),
            "elements": total_records,
            "tables": tables,
            "blocks": blocks,
            "json_data": None,
            "metadata": {
                **self._base_metadata(path),
                "title": path.stem,
                "sheet_count": len(tables),
                "record_count": total_records,
                "schema": schema,
            },
        }

    def _parse_pdf(self, path: Path) -> dict[str, Any]:
        warnings: list[str] = []
        pages: list[str] = []

        if PdfReader is not None:
            reader = PdfReader(str(path))
            for page in reader.pages:
                pages.append((page.extract_text() or "").strip())
        else:
            warnings.append("pypdf 未安装，PDF 将尝试通过 unstructured 降级读取。")
            pages = self._partition_pdf(path)

        text = "\n\n".join(page for page in pages if page).strip()
        blocks = [
            {
                "id": f"page_{index}",
                "type": "page",
                "text": page,
                "source_anchor": {"page_index": index},
            }
            for index, page in enumerate(pages, start=1)
            if page
        ]
        return {
            "type": "pdf",
            "content": text,
            "elements": len([page for page in pages if page]),
            "tables": [],
            "blocks": blocks,
            "json_data": None,
            "metadata": {
                **self._base_metadata(path),
                "title": self._first_non_empty_line(text) or path.stem,
                "page_count": len(pages),
                "text_length": len(text),
            },
            "warnings": warnings,
        }

    def _parse_image(self, path: Path) -> dict[str, Any]:
        warnings: list[str] = []
        metadata = self._base_metadata(path)
        text = ""

        try:
            from PIL import Image

            with Image.open(path) as image:
                metadata["image_size"] = {"width": image.width, "height": image.height}
        except Exception as exc:
            warnings.append(f"无法读取图片尺寸: {exc}")

        extracted = self._partition_image(path)
        if extracted:
            text = "\n".join(extracted).strip()

        blocks = [
            {
                "id": f"line_{index}",
                "type": "ocr_line",
                "text": line,
                "source_anchor": {"line_index": index},
            }
            for index, line in enumerate(extracted, start=1)
            if line
        ]

        return {
            "type": "image",
            "content": text,
            "elements": max(1, len(blocks)),
            "tables": [],
            "blocks": blocks,
            "json_data": None,
            "metadata": {
                **metadata,
                "title": path.stem,
                "text_length": len(text),
            },
            "warnings": warnings,
        }

    def _partition_docx(self, path: Path) -> list[str]:
        try:
            from unstructured.partition.docx import partition_docx

            return [str(element).strip() for element in partition_docx(filename=str(path)) if str(element).strip()]
        except Exception:
            return []

    def _partition_pdf(self, path: Path) -> list[str]:
        try:
            from unstructured.partition.pdf import partition_pdf

            return [str(element).strip() for element in partition_pdf(filename=str(path)) if str(element).strip()]
        except Exception:
            return []

    def _partition_image(self, path: Path) -> list[str]:
        try:
            from unstructured.partition.image import partition_image

            return [str(element).strip() for element in partition_image(filename=str(path)) if str(element).strip()]
        except Exception:
            return []

    def _build_table_payload(
        self,
        name: str,
        rows: list[list[str]],
        source_anchor: dict[str, Any] | None = None,
    ) -> dict[str, Any] | None:
        cleaned_rows = []
        for row in rows:
            normalized = self._trim_trailing_empty([self._clean_text(cell) for cell in row])
            if any(cell for cell in normalized):
                cleaned_rows.append(normalized)

        if not cleaned_rows:
            return None

        width = max(len(row) for row in cleaned_rows)
        padded_rows = [row + [""] * (width - len(row)) for row in cleaned_rows]

        headers = self._sanitize_headers(padded_rows[0])
        records = []
        record_sources = []
        for row_index, row in enumerate(padded_rows[1:], start=2):
            record = {header: row[index] for index, header in enumerate(headers)}
            if any(str(value).strip() for value in record.values()):
                records.append(record)
                record_sources.append({**(source_anchor or {}), "row_index": row_index})

        return {
            "name": name,
            "headers": headers,
            "records": records,
            "row_count": len(records),
            "column_count": len(headers),
            "preview_rows": padded_rows[: TABLE_PREVIEW_ROWS + 1],
            "source_anchor": source_anchor or {},
            "record_sources": record_sources,
        }

    def _sanitize_headers(self, headers: list[str]) -> list[str]:
        counts: dict[str, int] = {}
        normalized_headers = []

        for index, header in enumerate(headers, start=1):
            candidate = self._clean_text(header) or f"字段_{index}"
            if candidate in counts:
                counts[candidate] += 1
                candidate = f"{candidate}_{counts[candidate]}"
            else:
                counts[candidate] = 1
            normalized_headers.append(candidate)

        return normalized_headers

    def _tables_to_preview_text(self, tables: list[dict[str, Any]]) -> str:
        lines = []
        for table in tables:
            lines.append(f"[{table['name']}]")
            for row in table.get("preview_rows", []):
                lines.append("\t".join(row))
            lines.append("")

        preview = "\n".join(lines).strip()
        return preview[:TEXT_PREVIEW_CHARS]

    def _trim_trailing_empty(self, row: list[str]) -> list[str]:
        trimmed = list(row)
        while trimmed and not trimmed[-1]:
            trimmed.pop()
        return trimmed

    def _clean_text(self, value: Any) -> str:
        if value is None:
            return ""
        text = str(value).replace("\u3000", " ").strip()
        return text

    def _base_metadata(self, path: Path) -> dict[str, Any]:
        return {
            "file_name": path.name,
            "file_path": str(path),
            "file_size": path.stat().st_size if path.exists() else 0,
        }

    def _first_non_empty_line(self, text: str) -> str:
        for line in text.splitlines():
            stripped = line.strip()
            if stripped:
                return stripped
        return ""
